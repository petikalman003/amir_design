"""Real BLAST -> model -> R tests on deliberately synthetic transcript references."""

from io import BytesIO
from pathlib import Path

import numpy as np
import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

from amir_design.annotations import prepare_reference_annotations
from amir_design.offtargets import SeedMatchRunner
from amir_design.pipeline import run_identifier_pipeline, run_pipeline
from amir_design.references import prepare_reference_bundle
from amir_design.reports import workbook_bytes
from amir_design.scoring import ShRNAIScorer


ROOT = Path(__file__).resolve().parents[1]
TARGET = "CGGCTGCGCATATACATATATA"


@pytest.fixture(scope="module")
def real_reference(tmp_path_factory):
    root = tmp_path_factory.mktemp("integration") / "human"
    return prepare_reference_bundle(
        root, species="human", source_fasta=ROOT / "tests/fixtures/integration_transcripts.fa",
        release="synthetic-fixture", assembly="synthetic-fixture",
    )


@pytest.fixture(scope="module")
def annotated_reference(tmp_path_factory):
    root = tmp_path_factory.mktemp("annotated-integration") / "human"
    reference = prepare_reference_bundle(
        root, species="human", source_fasta=ROOT / "tests/fixtures/integration_transcripts.fa",
        release="synthetic-fixture", assembly="synthetic-fixture",
    )
    records = [
        ("TEST_GENE", "TEST_TX1", "exon", 100, 127),
        ("TEST_GENE", "TEST_TX1", "CDS", 102, 125),
        ("TEST_GENE", "TEST_TX1", "UTR", 100, 101),
        ("TEST_GENE", "TEST_TX1", "UTR", 126, 127),
        ("TEST_GENE", "TEST_TX2", "exon", 300, 325),
        ("TEST_GENE", "TEST_TX2", "CDS", 301, 324),
        ("TEST_GENE", "TEST_TX2", "UTR", 300, 300),
        ("TEST_GENE", "TEST_TX2", "UTR", 325, 325),
        ("TEST_OFF_GENE1", "TEST_OFF_TX1", "exon", 1000, 1015),
        ("TEST_OFF_GENE1", "TEST_OFF_TX2", "exon", 1100, 1115),
        ("TEST_OFF_GENE2", "TEST_OFF_TX3", "exon", 2000, 2017),
    ]
    gtf = root.parent / "regions.gtf"
    gtf.write_text("".join(
        f'chr1\tfixture\t{feature}\t{start}\t{end}\t.\t+\t.\t'
        f'gene_id "{gene}"; transcript_id "{transcript}";\n'
        for gene, transcript, feature, start, end in records
    ))
    prepare_reference_annotations(reference, source_gtf=gtf)
    return reference


@pytest.fixture(scope="module")
def real_dependencies():
    scorer = ShRNAIScorer(ROOT / "tools/shrnai")
    runner = SeedMatchRunner(ROOT / "tools/seedmatchr", rscript="/workspace/.runtime/bin/Rscript")
    return scorer, runner


def test_real_three_stage_pipeline_and_download(real_reference, real_dependencies):
    result = run_pipeline(TARGET, real_reference, *real_dependencies)
    assert result.status == "complete"
    assert len(result.summary) == 1
    row = result.summary.iloc[0]
    assert np.isfinite(row.predicted_efficiency_score)
    assert row.target_isoforms_covered == row.target_isoforms_total == 2
    assert row.off_target_genes == 2
    assert row.off_target_transcripts == row.off_target_sites == 3
    assert row.mer8_sites == 3
    assert set(result.sites.gene_id) == {"TEST_OFF_GENE1", "TEST_OFF_GENE2"}
    assert result.metadata["synthesis_status"] == "insert_available"
    assert len(row.mire_insert_dna_5to3) == 119
    assert row.mire_insert_dna_5to3[68:90] == row.guide_dna_5to3
    assert result.metadata["species_reference"]["source"]["release"] == "synthetic-fixture"
    report = openpyxl.load_workbook(BytesIO(workbook_bytes(result)))
    assert report["Candidates"].max_row == 2
    assert report["Off-target sites"].max_row == 4
    assert report["Target isoform sites"].max_row == 3


def test_real_reference_without_shared_isoform_returns_no_candidate(tmp_path):
    reference = prepare_reference_bundle(
        tmp_path / "no-shared", species="mouse",
        source_fasta=ROOT / "tests/fixtures/no_shared_isoform_transcripts.fa",
        release="synthetic-fixture", assembly="synthetic-fixture",
    )
    result = run_pipeline(TARGET, reference, object(), object())
    assert result.status == "no_eligible_candidate"
    assert result.summary.empty


def test_streamlit_real_submission_and_failed_submission_clear_old_result(
    real_reference, monkeypatch,
):
    calls = {"score": 0, "screen": 0}
    original_score = ShRNAIScorer.score
    original_screen = SeedMatchRunner.screen

    def tracked_score(self, *args, **kwargs):
        calls["score"] += 1
        return original_score(self, *args, **kwargs)

    def tracked_screen(self, *args, **kwargs):
        calls["screen"] += 1
        return original_screen(self, *args, **kwargs)

    monkeypatch.setattr(ShRNAIScorer, "score", tracked_score)
    monkeypatch.setattr(SeedMatchRunner, "screen", tracked_screen)
    monkeypatch.setenv("AMIR_REFERENCE_ROOT", str(real_reference.root.parent))
    monkeypatch.setenv("AMIR_TOOLS_ROOT", str(ROOT / "tools"))
    monkeypatch.setenv("AMIR_RSCRIPT", "/workspace/.runtime/bin/Rscript")
    monkeypatch.delenv("AMIR_MIRE_TEMPLATE", raising=False)
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
    assert not app.exception
    app.text_input[0].set_value("TEST_GENE")
    app.selectbox[0].select("Whole transcript")
    app.button[0].click().run()
    assert not app.exception and not app.error
    summary = app.session_state["result"].summary
    assert len(summary) == 4
    shared_row = summary.loc[summary.target_site_dna_5to3 == TARGET].iloc[0]
    assert shared_row.off_target_genes == 2
    assert set(summary.target_isoforms_covered) == {1, 2}
    assert len(app.session_state["workbook"]) > 1000
    workbook = openpyxl.load_workbook(BytesIO(app.session_state["workbook"]))
    assert "Candidate origins" in workbook.sheetnames
    assert workbook["Candidate origins"].max_row == 6
    app.selectbox[-1].select("Isoform coverage").run()
    assert not app.exception and not app.error
    assert calls == {"score": 1, "screen": 1}
    assert app.session_state["workbook_sort"] == "coverage"
    sorted_workbook = openpyxl.load_workbook(BytesIO(app.session_state["workbook"]))
    run_information = dict(sorted_workbook["Run information"].iter_rows(min_row=2, values_only=True))
    assert "coverage" in run_information["sort_order"].lower()
    assert sorted_workbook["Candidates"]["A2"].value == shared_row.candidate_id
    app.checkbox[0].check()
    app.button[0].click().run()
    assert not app.exception and not app.error
    assert app.session_state["result"].metadata["require_all_isoforms"] is True
    assert len(app.session_state["result"].summary) == 1
    assert app.session_state["result"].summary.iloc[0].target_isoforms_covered == 2
    app.text_input[0].set_value("INVALID_GENE_ID")
    app.button[0].click().run()
    assert not app.exception
    assert len(app.error) == 1
    assert "result" not in app.session_state


def test_identifier_gene_and_transcript_keep_origin_and_full_gene_coverage(
    real_reference, real_dependencies,
):
    gene_result = run_identifier_pipeline("TEST_GENE", real_reference, *real_dependencies, region="whole")
    transcript_result = run_identifier_pipeline(
        "TEST_TX1", real_reference, *real_dependencies, region="whole",
    )
    assert gene_result.status == transcript_result.status == "complete"
    gene_shared = gene_result.summary.loc[gene_result.summary.target_site_dna_5to3 == TARGET].iloc[0]
    transcript_shared = transcript_result.summary.loc[
        transcript_result.summary.target_site_dna_5to3 == TARGET
    ].iloc[0]
    assert gene_shared.target_isoforms_covered == 2
    assert transcript_shared.target_isoforms_covered == 2
    assert set(gene_result.summary.target_isoforms_covered) == {1, 2}
    assert set(gene_result.candidate_origins.transcript_id) == {"TEST_TX1", "TEST_TX2"}
    assert set(transcript_result.candidate_origins.transcript_id) == {"TEST_TX1"}
    assert set(transcript_result.on_target_sites.transcript_id) == {"TEST_TX1", "TEST_TX2"}


def test_streamlit_default_cds_and_utr_no_candidate(annotated_reference, monkeypatch):
    monkeypatch.setenv("AMIR_REFERENCE_ROOT", str(annotated_reference.root.parent))
    monkeypatch.setenv("AMIR_TOOLS_ROOT", str(ROOT / "tools"))
    monkeypatch.setenv("AMIR_RSCRIPT", "/workspace/.runtime/bin/Rscript")
    monkeypatch.delenv("AMIR_MIRE_TEMPLATE", raising=False)
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
    assert not app.exception
    assert app.selectbox[0].value == "CDS only"
    assert app.checkbox[0].value is False
    app.text_input[0].set_value("TEST_TX1")
    app.button[0].click().run()
    assert not app.exception and not app.error
    result = app.session_state["result"]
    assert result.metadata["search_region"] == "cds"
    assert result.metadata["input_identifier_kind"] == "transcript"
    shared = result.summary.loc[result.summary.target_site_dna_5to3 == TARGET].iloc[0]
    assert shared.target_isoforms_covered == 2
    assert set(result.candidate_origins.transcript_id) == {"TEST_TX1"}
    assert set(result.candidate_origins.region) == {"cds"}
    app.selectbox[0].select("3′ UTR only")
    app.button[0].click().run()
    assert not app.exception and not app.error
    assert app.session_state["result"].status == "no_eligible_candidate"
    assert app.session_state["result"].summary.empty
    assert len(app.warning) >= 1
