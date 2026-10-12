"""The single workbook retains large gene-wide fields without Excel truncation."""

from io import BytesIO
import json

import openpyxl
import pandas as pd

from amir_design.pipeline import PipelineResult
from amir_design.reports import workbook_bytes


def test_large_candidate_fields_are_preserved_in_ordered_literal_parts():
    positions = "=" + ";".join(f"TEST_TX{i}:4,5,6" for i in range(5000))
    result = PipelineResult(
        status="complete",
        summary=pd.DataFrame({
            "candidate_id": ["amiR_0001"],
            "source_transcript_start_positions_1based": [positions],
        }),
        sites=pd.DataFrame(),
        on_target_sites=pd.DataFrame(),
        metadata={"target_isoforms": [f"TEST_TX{i}" for i in range(5000)]},
    )
    workbook = openpyxl.load_workbook(BytesIO(workbook_bytes(result)))
    assert "Full text recorded in Long values" in workbook["Candidates"]["B2"].value
    values = workbook["Long values"]
    rows = list(values.iter_rows(min_row=2, values_only=True))
    assert "".join(row[-1] for row in rows) == positions
    assert {row[1] for row in rows} == {"Candidates"}
    assert {row[2] for row in rows} == {"B2"}
    assert [row[4] for row in rows] == list(range(1, len(rows) + 1))
    assert values["G2"].data_type == "s"
    metadata = list(workbook["Run information"].iter_rows(min_row=2, values_only=True))
    isoform_parts = [value for key, value in metadata if key.startswith("target_isoforms")]
    assert json.loads("".join(isoform_parts)) == result.metadata["target_isoforms"]
