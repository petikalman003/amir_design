from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pandas as pd
import pytest

from amir_design.candidates import reverse_complement
from amir_design.pipeline import (
    PipelineResult, run_identifier_pipeline, run_pipeline, sort_result, with_site_details,
)
from amir_design.reports import workbook_bytes


TARGET = "CGGCTGCGCATATACATATATA"


class ReferenceFixture:
    manifest = {"source": {"species": "test", "release": "fixture"}}
    fasta_path = Path("unused-fixture.fa")

    def __init__(self, missing_isoform=False):
        self.missing_isoform = missing_isoform

    def gene_for_cds(self, cds):
        assert cds == TARGET
        return {"gene_id": "target", "gene_name": "FixtureTarget"}

    def transcripts_for_gene(self, gene):
        return {"tx1": TARGET, "tx2": "ATG" * 30 if self.missing_isoform else "ATG" + TARGET}

    def transcript_gene_map(self):
        return {"tx1": "target", "tx2": "target", "off1": "other"}


class ScorerFixture:
    provenance = {"kind": "test stub"}

    def score(self, targets):
        assert targets == [TARGET]
        return [0.8]


class RunnerFixture:
    def screen(self, guides, fasta, gene, mapping, *, progress=None, include_sites=True):
        assert guides == {"amiR_0001": reverse_complement(TARGET)}
        assert gene == "target"
        if progress:
            progress("Screening guide 1/1")
        return SimpleNamespace(
            summary={"amiR_0001": {"off_target_genes": 1, "off_target_transcripts": 1,
                                    "off_target_sites": 1}},
            sites=([{"candidate_id": "amiR_0001", "gene_id": "other", "transcript_id": "off1"}]
                   if include_sites else []),
            provenance={"kind": "test stub"},
        )


def test_no_eligible_candidate_does_not_score_or_screen(tmp_path):
    result = run_pipeline(
        TARGET, ReferenceFixture(missing_isoform=True), object(), object(),
        template_path=tmp_path / "missing-template.json",
    )
    assert result.status == "no_eligible_candidate"
    assert result.summary.empty
    assert len(result.metadata["target_isoforms"]) == 2
    assert result.metadata["rejected_incomplete_isoform_coverage"] == 1
    workbook = openpyxl.load_workbook(BytesIO(workbook_bytes(result)))
    assert "Candidates" in workbook.sheetnames
    assert workbook["Candidates"].max_row == 1


def test_missing_template_preserves_real_stage_results_without_claiming_insert(monkeypatch):
    import amir_design.scaffold as scaffold

    monkeypatch.setattr(scaffold, "DEFAULT_MIRE_TEMPLATE", None)
    result = run_pipeline(TARGET, ReferenceFixture(), ScorerFixture(), RunnerFixture())
    assert result.summary.iloc[0].predicted_efficiency_score == 0.8
    assert result.summary.iloc[0].off_target_genes == 1
    assert result.summary.iloc[0].mire_insert_dna_5to3 == ""
    assert result.metadata["synthesis_status"] == "template_required"
    assert len(result.on_target_sites) == 2


def test_published_default_insert_contains_exact_scored_and_screened_guide():
    result = run_pipeline(TARGET, ReferenceFixture(), ScorerFixture(), RunnerFixture())
    row = result.summary.iloc[0]
    assert result.status == "complete"
    assert len(row.mire_insert_dna_5to3) == 119
    assert row.mire_insert_dna_5to3[68:90] == row.guide_dna_5to3
    assert result.metadata["synthesis_scaffold"]["template_origin"] == "verified published miR-E fragment"


def test_report_records_insert_provenance_and_literal_external_strings(monkeypatch):
    import amir_design.reports as reports

    result = run_pipeline(
        TARGET, ReferenceFixture(), ScorerFixture(), RunnerFixture(),
        insert_builder=lambda guide: "ACGT" + guide,
        scaffold_provenance={"source": "test fixture only"},
    )
    result.metadata["target_gene_name"] = "=SOME_EXTERNAL_IDENTIFIER"
    result.sites = pd.DataFrame({"gene_id": ["=literal", "+literal", "@literal", "-literal"]})
    result.summary.loc[:, "off_target_sites"] = 4
    monkeypatch.setattr(reports, "MAX_DATA_ROWS", 2)
    workbook = openpyxl.load_workbook(BytesIO(workbook_bytes(result)))
    assert "Off-target sites 2" in workbook.sheetnames
    cell = workbook["Off-target sites 1"]["A2"]
    assert cell.value == "=literal" and cell.data_type == "s"
    assert result.metadata["synthesis_scaffold"]["source"] == "test fixture only"


def test_stage_failure_is_not_reported_as_zero_offtargets():
    class BrokenRunner:
        def screen(self, *args, **kwargs):
            raise RuntimeError("R search failed")

    with pytest.raises(RuntimeError, match="R search failed"):
        run_pipeline(TARGET, ReferenceFixture(), ScorerFixture(), BrokenRunner())


def test_malformed_insert_is_not_treated_as_an_unavailable_template():
    def builder(guide):
        raise ValueError("invalid template")

    with pytest.raises(ValueError, match="invalid template"):
        run_pipeline(TARGET, ReferenceFixture(), ScorerFixture(), RunnerFixture(), insert_builder=builder)


class IdentifierReferenceFixture(ReferenceFixture):
    annotation_provenance = {"source": "matching fixture GTF"}

    def resolve_identifier(self, identifier):
        assert identifier in {"target", "tx1"}
        return {"gene_id": "target", "gene_name": "FixtureTarget",
                "input_type": "transcript" if identifier == "tx1" else "gene",
                "transcript_ids": ["tx1", "tx2"], "matching_transcript_ids": ["tx1"]}

    def regions_for_gene(self, gene):
        return {"tx1": {"cds": [(1, 22)], "utr3": [], "utr5": []},
                "tx2": {"cds": [], "utr3": [], "utr5": []}}


def test_identifier_defaults_keep_partial_coverage_with_gene_denominator_and_origins():
    result = run_identifier_pipeline(
        "tx1", IdentifierReferenceFixture(missing_isoform=True), ScorerFixture(), RunnerFixture(),
    )
    row = result.summary.iloc[0]
    assert row.target_isoforms_covered == 1
    assert row.target_isoforms_total == 2
    assert row.target_isoform_coverage_fraction == 0.5
    assert row.covered_transcript_ids == "tx1"
    assert row.missed_transcript_ids == "tx2"
    assert row.source_transcript_ids == "tx1"
    assert row.source_region == "cds"
    assert row.source_transcript_start_positions_1based == "tx1:1"
    assert result.metadata["source_transcript_ids"] == ["tx1"]
    assert result.metadata["input_identifier_kind"] == "transcript"
    assert result.metadata["require_all_isoforms"] is False
    assert result.metadata["annotation_provenance"] == IdentifierReferenceFixture.annotation_provenance
    assert len(result.on_target_sites) == 1
    origin = result.candidate_origins.iloc[0]
    assert (origin.transcript_id, origin.start_1based, origin.end_1based) == ("tx1", 1, 22)
    assert origin.target_site_dna_5to3 == TARGET


def test_checkbox_complete_coverage_returns_valid_empty_result_without_running_external_stages():
    result = run_identifier_pipeline(
        "target", IdentifierReferenceFixture(missing_isoform=True), object(), object(),
        require_all_isoforms=True,
    )
    assert result.status == "no_eligible_candidate"
    assert result.metadata["require_all_isoforms"] is True
    assert "every target isoform" in result.notices[0]
    assert result.metadata["rejected_incomplete_isoform_coverage"] == 1


def test_missing_selected_region_is_explicit_and_not_claimed_as_coverage_rejection():
    result = run_identifier_pipeline(
        "target", IdentifierReferenceFixture(), object(), object(), region="utr3",
    )
    assert result.status == "no_eligible_candidate"
    assert result.metadata["source_transcript_ids"] == ["tx1", "tx2"]
    assert result.metadata["missing_region_transcript_ids"] == ["tx1", "tx2"]
    assert result.metadata["rejected_incomplete_isoform_coverage"] == 0
    assert "no annotated sequence in this region" in result.notices[0]


def test_whole_transcript_design_never_requires_or_guesses_coding_annotations():
    class UnannotatedReference(IdentifierReferenceFixture):
        def regions_for_gene(self, gene):
            raise AssertionError("Whole-transcript search must not need region annotations")

    result = run_identifier_pipeline(
        "tx1", UnannotatedReference(), ScorerFixture(), RunnerFixture(), region="whole",
    )
    assert result.summary.iloc[0].target_isoforms_covered == 2
    assert result.metadata["search_region"] == "whole"
    assert result.metadata["missing_region_transcript_ids"] == []


def test_gene_sources_include_each_isoform_but_transcript_sources_only_selected_one():
    other = "CCGCTGCGCATATACATATATA"

    class DifferentIsoforms(IdentifierReferenceFixture):
        def transcripts_for_gene(self, gene):
            return {"tx1": TARGET, "tx2": other}

        def regions_for_gene(self, gene):
            return {"tx1": {"cds": [(1, 22)]}, "tx2": {"cds": [(1, 22)]}}

    class Scorer:
        provenance = {"kind": "test stub"}

        def score(self, targets):
            return [0.8] * len(targets)

    class Runner:
        def screen(self, guides, fasta, gene, mapping, *, progress=None, include_sites=True):
            assert gene == "target"
            return SimpleNamespace(
                summary={candidate: {"off_target_genes": 0, "off_target_transcripts": 0,
                                     "off_target_sites": 0, "mer8_sites": 0}
                         for candidate in guides},
                sites=[], provenance={"kind": "test stub"},
            )

    gene_result = run_identifier_pipeline("target", DifferentIsoforms(), Scorer(), Runner())
    assert set(gene_result.summary.target_site_dna_5to3) == {TARGET, other}
    assert set(gene_result.summary.source_transcript_ids) == {"tx1", "tx2"}
    assert list(gene_result.summary.target_isoforms_covered) == [1, 1]
    assert list(gene_result.summary.target_isoforms_total) == [2, 2]
    assert "mer8_sites" in gene_result.summary
    transcript_result = run_identifier_pipeline("tx1", DifferentIsoforms(), Scorer(), Runner())
    assert list(transcript_result.summary.target_site_dna_5to3) == [TARGET]
    assert transcript_result.summary.iloc[0].target_isoforms_total == 2


def test_pipeline_reports_stage_timing_and_forwards_per_guide_progress():
    messages = []
    result = run_pipeline(
        TARGET, ReferenceFixture(), ScorerFixture(), RunnerFixture(), progress=messages.append,
    )
    assert "Screening guide 1/1" in messages
    assert messages[-1] == "Assembling candidate sequences and isoform coverage results"
    assert set(result.metadata["stage_duration_seconds"]) == {
        "candidate_selection", "shrnai", "seedmatchr", "result_assembly",
    }
    assert result.metadata["analysis_duration_seconds"] >= 0


def test_optional_site_positions_do_not_change_counts_or_repeat_scoring():
    reference = ReferenceFixture()
    result = run_pipeline(
        TARGET, reference, ScorerFixture(), RunnerFixture(), include_site_details=False,
    )
    assert result.sites.empty
    assert result.summary.iloc[0].off_target_sites == 1
    assert result.metadata["off_target_site_details_retained"] is False
    detailed = with_site_details(result, reference, RunnerFixture())
    assert len(detailed.sites) == 1
    pd.testing.assert_frame_equal(detailed.summary, result.summary)
    pd.testing.assert_frame_equal(detailed.on_target_sites, result.on_target_sites)
    assert detailed.metadata["detailed_screen_counts_verified"] is True
    assert with_site_details(detailed, object(), object()) is detailed


def test_detailed_export_rejects_inconsistent_counter_instead_of_changing_results():
    reference = ReferenceFixture()
    result = run_pipeline(
        TARGET, reference, ScorerFixture(), RunnerFixture(), include_site_details=False,
    )

    class ChangedRunner(RunnerFixture):
        def screen(self, *args, **kwargs):
            response = super().screen(*args, **kwargs)
            response.summary["amiR_0001"]["off_target_genes"] += 1
            return response

    with pytest.raises(RuntimeError, match="changed off_target_genes"):
        with_site_details(result, reference, ChangedRunner())
    assert result.metadata["off_target_site_details_retained"] is False
    assert result.sites.empty


def test_sorting_changes_independent_order_and_report_without_changing_measurements():
    summary = pd.DataFrame([
        {"candidate_id": "high_activity", "predicted_efficiency_score": 0.9,
         "off_target_genes": 10, "target_isoform_coverage_fraction": 0.5},
        {"candidate_id": "high_coverage", "predicted_efficiency_score": 0.8,
         "off_target_genes": 5, "target_isoform_coverage_fraction": 1.0},
        {"candidate_id": "low_offtargets", "predicted_efficiency_score": 0.7,
         "off_target_genes": 0, "target_isoform_coverage_fraction": 0.5},
    ])
    result = PipelineResult("complete", summary, pd.DataFrame(), pd.DataFrame(), {"sort_by": "old"})
    for order, first in [("activity", "high_activity"), ("coverage", "high_coverage"),
                         ("off_targets", "low_offtargets")]:
        sorted_result = sort_result(result, order)
        assert sorted_result.summary.iloc[0].candidate_id == first
        assert sorted_result.metadata["sort_by"] == order
        pd.testing.assert_frame_equal(
            sorted_result.summary.sort_values("candidate_id").reset_index(drop=True),
            summary.sort_values("candidate_id").reset_index(drop=True),
        )
    assert result.metadata["sort_by"] == "old"
    assert result.summary.iloc[0].candidate_id == "high_activity"
    with pytest.raises(ValueError, match="Unknown candidate sort"):
        sort_result(result, "combined")
